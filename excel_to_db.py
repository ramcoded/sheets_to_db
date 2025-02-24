import openpyxl
import csv
import os
import re

def sanitize_column_name(name, index=None):
    # Replace invalid characters with underscores and handle unnamed columns
    if name is None or name.strip() == '':
        return f"unnamed_column_{index}" if index is not None else 'unnamed_column'
    return re.sub(r'\W|^(?=\d)', '_', name)

def preprocess_worksheet(ws, skip_columns=None, skip_rows=None):
    if skip_columns is None:
        skip_columns = []  # Default list to skip specific columns by their indices
    if skip_rows is None:
        skip_rows = []  # Default list to skip specific rows by their indices

    print(f"Preprocessing worksheet with skip_columns={skip_columns} and skip_rows={skip_rows}")

    # Read the data rows, skipping the specified rows and columns
    data = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx in skip_rows:
            continue
        filtered_row = [value for i, value in enumerate(row) if i not in skip_columns]
        data.append(filtered_row)
    
    return data

def preprocess_csv(file_path, skip_columns=None, skip_rows=None):
    if skip_columns is None:
        skip_columns = []  # Default list to skip specific columns by their indices
    if skip_rows is None:
        skip_rows = []  # Default list to skip specific rows by their indices

    print(f"Preprocessing CSV with skip_columns={skip_columns} and skip_rows={skip_rows}")

    with open(file_path, newline='', encoding='utf-8') as csvfile:
        reader = csv.reader(csvfile)
        data = []
        for row_idx, row in enumerate(reader):
            if row_idx in skip_rows:
                continue
            filtered_row = [value for i, value in enumerate(row) if i not in skip_columns]
            data.append(filtered_row)
        
        return data

def convert_excel_to_db(file_path, db_connection, skip_columns=None, skip_rows=None):
    # Determine the file extension
    file_extension = os.path.splitext(file_path)[1]

    # Read the file based on its extension
    if file_extension == '.xlsx' or file_extension == '.xls':
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        records = preprocess_worksheet(ws, skip_columns, skip_rows)
    elif file_extension == '.csv':
        records = preprocess_csv(file_path, skip_columns, skip_rows)
    else:
        raise ValueError("Unsupported file format. Please provide an Excel or CSV file.")

    cursor = db_connection.cursor()

    # Define the fixed columns for the database
    fixed_columns = [
        'Location', 'important', 'january', 'february', 'march', 'april', 'may', 'june',
        'july', 'august', 'september', 'october', 'november', 'december', 'total'
    ]

    # Insert data into the table
    for record in records:
        # Prepare the SQL insert statement
        columns = ', '.join([f"`{col}`" for col in fixed_columns])
        placeholders = ', '.join(['%s'] * len(fixed_columns))
        sql = f"INSERT INTO data_table ({columns}) VALUES ({placeholders})"
        
        # Map the record values to the specific columns
        values = [record[i] if i < len(record) else '' for i in range(len(fixed_columns))]
        
        # Execute the insert statement
        cursor.execute(sql, tuple(values))

    # Commit the transaction
    db_connection.commit()
    cursor.close()
